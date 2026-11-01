#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import unittest
import os
from outils.generer_facture_word import generer_facture_word
from outils.generer_rapport_excel import creer_rapport_comptable_excel
import docx
import openpyxl

class TestOutilsOffice(unittest.TestCase):
    def setUp(self):
        self.doc_path = "test_facture.docx"
        self.xls_path = "test_rapport.xlsx"

    def tearDown(self):
        if os.path.exists(self.doc_path):
            os.remove(self.doc_path)
        if os.path.exists(self.xls_path):
            os.remove(self.xls_path)

    def test_generation_facture_word(self):
        res = generer_facture_word(output_path=self.doc_path)
        self.assertTrue(os.path.exists(self.doc_path))
        doc = docx.Document(self.doc_path)
        # Vérifier le texte dans les paragraphes et les tables
        texte_paragraphes = " ".join([p.text for p in doc.paragraphs])
        texte_tables = " ".join([cell.text for table in doc.tables for row in table.rows for cell in row.cells])
        texte_global = texte_paragraphes + " " + texte_tables
        self.assertIn("18 EITHING SARL", texte_global)
        self.assertIn("CODE DE LIVRAISON CLIENT", texte_global)
        self.assertIn("4892", texte_global)

    def test_generation_rapport_excel(self):
        res = creer_rapport_comptable_excel(output_path=self.xls_path)
        self.assertTrue(os.path.exists(self.xls_path))
        wb = openpyxl.load_workbook(self.xls_path)
        self.assertIn("Journal Ventes & Marges", wb.sheetnames)
        self.assertIn("Réconciliation MoMo & Banques", wb.sheetnames)
        self.assertIn("Commissions Livreurs", wb.sheetnames)
        ws_livreurs = wb["Commissions Livreurs"]
        # Vérifier que le taux unitaire de 500 F est bien présent
        self.assertEqual(ws_livreurs["E4"].value, 500)

if __name__ == "__main__":
    unittest.main()
