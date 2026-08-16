import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def main():
    doc = docx.Document()

    # Set page margins (1 inch = 1440 dxa)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DOSSIER STRATÉGIQUE D'INVESTISSEMENT\nET PLAN D'AFFAIRES INDUSTRIEL")
    run.font.name = 'Arial'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0E, 0x2C, 0x43) # Navy blue

    # Subtitle / Header info block
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub.add_run("ENTREPRISE : 18 EITHING SARL\nKétou, Département du Plateau, République du Bénin\nSecteur : Recyclage, Broyage et Valorisation des Déchets Plastiques (PP & PEHD)")
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph() # Spacer

    with open("DOSSIER_STRATEGIQUE_18_EITHING.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def process_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for line in t_lines:
            if line.strip().startswith('|') and '---' not in line:
                parts = [p.strip().replace('**', '') for p in line.strip().split('|')[1:-1]]
                rows_data.append(parts)

        if not rows_data:
            return

        num_cols = len(rows_data[0])
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row in enumerate(rows_data):
            for c_idx, val in enumerate(row):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = val
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

                    # Formatting Header vs Data
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.name = 'Calibri'
                        run.font.size = Pt(10)
                        if r_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                    if r_idx == 0:
                        set_cell_background(cell, "0E2C43") # Dark Navy Header
                    elif r_idx % 2 == 1:
                        set_cell_background(cell, "F4EFE9") # Light Beige zebra stripe
                    else:
                        set_cell_background(cell, "FFFFFF")

        doc.add_paragraph() # Spacer after table

    for line in lines:
        raw_line = line.strip()

        # Skip main Title/Subtitle markdown lines as they are handled above
        if raw_line.startswith("# DOSSIER STRATÉGIQUE") or raw_line.startswith("**ENTREPRISE :**") or raw_line.startswith("**SIÈGE ET SITE") or raw_line.startswith("**SECTEUR D'ACTIVITÉ :**"):
            continue

        if raw_line.startswith('|'):
            in_table = True
            table_lines.append(raw_line)
            continue
        elif in_table:
            in_table = False
            process_table(table_lines)
            table_lines = []

        if raw_line.startswith('## '):
            h = doc.add_heading(level=1)
            run = h.add_run(raw_line.replace('## ', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0E, 0x2C, 0x43)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)

        elif raw_line.startswith('### '):
            h = doc.add_heading(level=2)
            run = h.add_run(raw_line.replace('### ', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xB6, 0x6A, 0x3F) # Accent warm brown
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)

        elif raw_line.startswith('#### '):
            h = doc.add_heading(level=3)
            run = h.add_run(raw_line.replace('#### ', ''))
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x2D, 0x3D)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)

        elif raw_line.startswith('* ') or raw_line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            text = raw_line[2:]
            # Handle bold formatting inside bullet
            parts = text.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.bold = True

        elif raw_line and raw_line[0].isdigit() and raw_line[1:3] in ['. ', ') ']:
            p = doc.add_paragraph(style='List Number')
            text = raw_line[3:]
            parts = text.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.bold = True

        elif raw_line == '---':
            continue

        elif raw_line:
            p = doc.add_paragraph()
            parts = raw_line.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.bold = True
            p.paragraph_format.space_after = Pt(6)

    if in_table:
        process_table(table_lines)

    doc.save("DOSSIER_STRATEGIQUE_18_EITHING.docx")
    print("Document Word DOSSIER_STRATEGIQUE_18_EITHING.docx généré avec succès!")

if __name__ == '__main__':
    main()
